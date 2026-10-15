"""Take Microsoft Store screenshots (1920x1080 PNG) of the real app with headless Edge.

Usage:  .venv\\Scripts\\python packaging\\screenshots.py [out_dir]       (default: store/screenshots)
Starts the Flask app on a free localhost port, makes a few sample files, then captures several views via
#view=<name>&notour (and ?open=[...] for the file list). Re-run after UI changes; output is overwritten.
"""
import json
import os
import shutil
import subprocess
import sys
import tempfile
import threading
import time
import urllib.parse
import urllib.request
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "store" / "screenshots"
EDGE = next((p for p in (os.path.expandvars(r"%ProgramFiles(x86)%\Microsoft\Edge\Application\msedge.exe"),
                         os.path.expandvars(r"%ProgramFiles%\Microsoft\Edge\Application\msedge.exe")) if os.path.exists(p)), None)
W, H = 1920, 1080


def samples(folder: Path) -> list[str]:
    folder.mkdir(parents=True, exist_ok=True)
    out = []
    from PIL import Image
    img = Image.open(ROOT / "web" / "assets" / "mascot-light.png").convert("RGBA")
    bg = Image.new("RGB", (1600, 1200), (238, 233, 255)); bg.paste(img.resize((600, 778)), (500, 211), img.resize((600, 778)))
    bg.save(folder / "Holiday photo.jpg", quality=92); out.append(folder / "Holiday photo.jpg")
    bg.save(folder / "Scan 2026-03.png"); out.append(folder / "Scan 2026-03.png")
    try:
        import pymupdf
        doc = pymupdf.open()
        for i in range(3):
            page = doc.new_page()
            page.insert_text((72, 96), f"Quarterly report - page {i + 1}", fontsize=22)
            page.insert_text((72, 140), "Converted, compressed and edited offline with Dr. Doc.", fontsize=12)
        doc.save(folder / "Quarterly report.pdf"); out.append(folder / "Quarterly report.pdf")
    except Exception as e:
        print("pdf sample skipped:", e)
    try:
        import docx
        d = docx.Document(); d.add_heading("Meeting notes", 1); d.add_paragraph("Agenda, decisions and next steps.")
        d.save(folder / "Meeting notes.docx"); out.append(folder / "Meeting notes.docx")
    except Exception as e:
        print("docx sample skipped:", e)
    try:
        import openpyxl
        wb = openpyxl.Workbook(); ws = wb.active
        ws.append(["Month", "Sales"]); [ws.append([m, 100 + i * 17]) for i, m in enumerate(["Jan", "Feb", "Mar", "Apr"])]
        wb.save(folder / "Budget 2026.xlsx"); out.append(folder / "Budget 2026.xlsx")
    except Exception as e:
        print("xlsx sample skipped:", e)
    from core.util import FFMPEG
    if FFMPEG:
        subprocess.run([FFMPEG, "-y", "-v", "error", "-f", "lavfi", "-i", "testsrc2=size=1280x720:rate=30", "-f", "lavfi",
                        "-i", "sine=frequency=440", "-t", "4", "-c:v", "libx264", "-pix_fmt", "yuv420p", "-c:a", "aac",
                        str(folder / "Birthday video.mp4")], creationflags=0x08000000)
        if (folder / "Birthday video.mp4").exists():
            out.append(folder / "Birthday video.mp4")
    return [str(p) for p in out]


def main():
    if not EDGE:
        sys.exit("Microsoft Edge not found")
    import app as drdoc
    from waitress import create_server
    port = drdoc.free_port(8890)
    server = create_server(drdoc.app, host="127.0.0.1", port=port, threads=8)
    threading.Thread(target=server.run, daemon=True).start()
    base = f"http://127.0.0.1:{port}/"
    for _ in range(50):
        try:
            urllib.request.urlopen(base + "api/caps", timeout=2); break
        except Exception:
            time.sleep(0.2)

    work = Path(tempfile.mkdtemp(prefix="drdoc-shots-"))
    files = samples(work / "Sample files")
    shots = [
        ("01-convert-files", base + "?open=" + urllib.parse.quote(json.dumps(files)) + "#notour", False),
        ("02-home", base + "#view=convert&notour", False),
        ("03-compress", base + "#view=" + urllib.parse.quote("tools:Compress") + "&notour", False),
        ("04-pdf-tools", base + "#view=" + urllib.parse.quote("tools:PDF") + "&notour", False),
        ("05-docs-and-sheets", base + "#view=office&notour", False),
        ("06-video-editor", base + "#view=video&notour", False),
        ("07-clean-up-night", base + "#view=cleanup&notour", True),
    ]
    OUT.mkdir(parents=True, exist_ok=True)
    for old in OUT.glob("*.png"):
        old.unlink()
    for name, url, dark in shots:
        profile = work / f"profile-{name}"
        png = OUT / f"{name}.png"
        args = [EDGE, "--headless=new", "--disable-gpu", "--hide-scrollbars", f"--window-size={W},{H}",
                f"--user-data-dir={profile}", "--no-first-run", "--virtual-time-budget=9000", "--run-all-compositor-stages-before-draw",
                f"--screenshot={png}"]
        if dark:
            args.append("--force-dark-mode")
            args.append("--blink-settings=preferredColorScheme=0")
        subprocess.run(args + [url], capture_output=True, timeout=120)
        print(("ok   " if png.exists() else "FAIL ") + str(png))
    server.close()
    shutil.rmtree(work, ignore_errors=True)


if __name__ == "__main__":
    main()
    sys.stdout.flush()
    os._exit(0)  # don't wait for the app's worker threads
