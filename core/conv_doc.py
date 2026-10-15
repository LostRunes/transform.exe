"""Document conversions: Word/Excel/PowerPoint via local MS Office (if installed),
plus pure-Python fallbacks for docx/md/txt/html."""
import html as htmllib
import sys
import threading
from pathlib import Path

from . import conv_pdf

_office_lock = threading.Lock()


def _office_available(prog_id: str) -> bool:
    if sys.platform != "win32":
        return False
    try:
        import winreg
        winreg.CloseKey(winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, prog_id))
        import win32com.client  # noqa: F401
        return True
    except Exception:
        return False


HAS_WORD = _office_available("Word.Application")
HAS_EXCEL = _office_available("Excel.Application")
HAS_POWERPOINT = _office_available("PowerPoint.Application")

# Word SaveAs2 format codes
WORD_FORMATS = {"pdf": 17, "docx": 16, "doc": 0, "rtf": 6, "odt": 23, "txt": 7, "html": 10, "xps": 18}
WORD_IN = {"doc", "docx", "docm", "dot", "dotx", "rtf", "odt", "wpd", "wps"}
EXCEL_IN = {"xls", "xlsx", "xlsm", "xlsb", "ods"}
PPT_IN = {"ppt", "pptx", "pptm", "pps", "ppsx", "odp"}


def _com(fn, src=None):
    """Run an Office automation call with COM initialised on this thread."""
    if src is not None and Path(src).stat().st_size == 0:
        raise ValueError("This file is empty")  # Office would happily export a blank (or 0-page) document
    import pythoncom
    import pywintypes
    with _office_lock:
        pythoncom.CoInitialize()
        try:
            return fn()
        except pywintypes.com_error as e:  # raw COM errors are tuples of HRESULTs and help-file paths
            info = e.excepinfo if isinstance(getattr(e, "excepinfo", None), tuple) else None
            detail = (info[2] if info and len(info) > 2 and info[2] else e.strerror) or "unknown error"
            raise RuntimeError(f"Office couldn't convert this file: {str(detail).splitlines()[0]} "
                               "(is it damaged, password-protected or open in another program?)") from None
        finally:
            pythoncom.CoUninitialize()


def _html_for_word(src: Path, tmp: Path) -> Path:
    """Word blocks forever (an invisible dialog) on a picture or stylesheet it can't load, e.g. a relative
    path that only made sense next to the original file, or a web URL while offline. Hand it a copy where
    local files are absolute and everything else is dropped."""
    import re
    from urllib.parse import unquote

    def local(url):
        if url.startswith("data:"):
            return url
        is_drive = re.match(r"^[a-z]:[\\/]", url, re.I)
        if re.match(r"^[a-z][a-z0-9+.-]*:", url, re.I) and not url.lower().startswith("file:") and not is_drive:
            return None  # http(s) & co.
        path = unquote(re.sub(r"^file:/*", "", url, flags=re.I).split("?")[0].split("#")[0])
        f = Path(path) if Path(path).is_absolute() else src.parent / path
        return f.resolve().as_uri() if f.is_file() else None

    def fix(m):
        tag = m.group(0)
        ref = re.search(r"""\b(src|href)\s*=\s*(["']?)([^"'\s>]+)\2""", tag, re.I)
        if not ref:
            return tag
        url = local(htmllib.unescape(ref.group(3)))
        return tag[:ref.start(3)] + htmllib.escape(url) + tag[ref.end(3):] if url else ""
    html = re.sub(r"(?is)<script\b[^>]*\bsrc\b[^>]*>.*?</script>", "", read_text(src))
    html = re.sub(r"(?i)<(img|link)\b[^>]*>", fix, html)
    tmp.write_text(html, encoding="utf-8")
    return tmp


def word_convert(src: Path, out: Path, params: dict):
    fmt = out.suffix.lower().lstrip(".")
    if src.suffix.lower() in (".html", ".htm") and not src.name.endswith(".word-src.html"):
        if src.stat().st_size == 0:
            raise ValueError("This file is empty")
        clean = _html_for_word(src, out.with_name(out.stem + ".word-src.html"))
        try:
            return word_convert(clean, out, params)
        finally:
            clean.unlink(missing_ok=True)

    def run():
        import win32com.client
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        try:
            doc = word.Documents.Open(str(src.resolve()), ConfirmConversions=False, ReadOnly=True,
                                      AddToRecentFiles=False, Visible=False)
            try:
                if fmt == "pdf":
                    doc.ExportAsFixedFormat(str(out.resolve()), 17, OptimizeFor=0)
                else:
                    if fmt in ("html", "txt"):
                        doc.WebOptions.Encoding = 65001  # UTF-8; Word's default is the ANSI code page
                    doc.SaveAs2(str(out.resolve()), FileFormat=WORD_FORMATS[fmt],
                                **({"Encoding": 65001} if fmt == "txt" else {}))
            finally:
                doc.Close(False)
        finally:
            word.Quit()
    _com(run, src)
    return [out]


def excel_to_pdf(src: Path, out: Path, params: dict):
    def run():
        import win32com.client
        xl = win32com.client.DispatchEx("Excel.Application")
        xl.Visible = False
        xl.DisplayAlerts = False
        try:
            wb = xl.Workbooks.Open(str(src.resolve()), ReadOnly=True, UpdateLinks=0)
            try:
                for ws in wb.Worksheets:  # fit each sheet to page width
                    try:
                        ws.PageSetup.Zoom = False
                        ws.PageSetup.FitToPagesWide = 1
                        ws.PageSetup.FitToPagesTall = False
                    except Exception:
                        pass
                wb.ExportAsFixedFormat(0, str(out.resolve()))
            finally:
                wb.Close(False)
        finally:
            xl.Quit()
    _com(run, src)
    return [out]


def ppt_to_pdf(src: Path, out: Path, params: dict):
    def run():
        import win32com.client
        pp = win32com.client.DispatchEx("PowerPoint.Application")
        try:
            pres = pp.Presentations.Open(str(src.resolve()), ReadOnly=True, Untitled=False, WithWindow=False)
            try:
                pres.SaveAs(str(out.resolve()), 32)  # ppSaveAsPDF
            finally:
                pres.Close()
        finally:
            pp.Quit()
    _com(run, src)
    return [out]


def ppt_to_pptx(src: Path, out: Path, params: dict):
    def run():
        import win32com.client
        pp = win32com.client.DispatchEx("PowerPoint.Application")
        try:
            pres = pp.Presentations.Open(str(src.resolve()), ReadOnly=True, Untitled=False, WithWindow=False)
            try:
                pres.SaveAs(str(out.resolve()), 24)  # ppSaveAsOpenXMLPresentation
            finally:
                pres.Close()
        finally:
            pp.Quit()
    _com(run, src)
    return [out]


# ---------------- pure-Python fallbacks ----------------

def read_text(path) -> str:
    """Text files in whatever Windows gave us: UTF-8 (with or without BOM), UTF-16 (Notepad's "Unicode"), or ANSI."""
    raw = Path(path).read_bytes()
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return raw.decode("utf-16", errors="replace")
    if raw.startswith(b"\xef\xbb\xbf"):
        raw = raw[3:]
    try:
        return raw.decode("utf-8")
    except UnicodeDecodeError:
        return raw.decode("cp1252", errors="replace")


def docx_to_html(src: Path, out: Path, params: dict):
    """Basic but dependable: headings, paragraphs (bold/italic), lists, tables."""
    import docx
    d = docx.Document(str(src))
    parts = []
    body = d.element.body
    for child in body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            p = docx.text.paragraph.Paragraph(child, d)
            runs = []
            for item in p.iter_inner_content():  # p.runs skips hyperlinks, which lost their text
                if isinstance(item, docx.text.hyperlink.Hyperlink):
                    inner = "".join(_run_html(r) for r in item.runs)
                    runs.append(f"<a href='{htmllib.escape(item.url)}'>{inner}</a>" if item.url else inner)
                else:
                    runs.append(_run_html(item))
            text = "".join(runs)
            style = (p.style.name or "").lower() if p.style is not None else ""
            if style.startswith("heading") and style[-1:].isdigit():
                n = min(int(style[-1]), 6)
                parts.append(f"<h{n}>{text}</h{n}>")
            elif style == "title":
                parts.append(f"<h1>{text}</h1>")
            elif "list" in style:
                tag = "ol" if "number" in style else "ul"
                parts.append(f"<{tag}><li>{text}</li></{tag}>")
            else:
                parts.append(f"<p>{text or '&nbsp;'}</p>")
        elif tag == "tbl":
            t = docx.table.Table(child, d)
            rows = []
            for row in t.rows:
                cells = "".join(f"<td{f' colspan={n}' if n > 1 else ''}>{htmllib.escape(c.text)}</td>"
                                for c, n in _row_cells(row))
                rows.append(f"<tr>{cells}</tr>")
            parts.append("<table>" + "".join(rows) + "</table>")
    html = "<!doctype html><html><head><meta charset='utf-8'></head><body>" + "\n".join(parts) + "</body></html>"
    out.write_text(html.replace("</ul>\n<ul>", "").replace("</ol>\n<ol>", ""), encoding="utf-8")
    return [out]


def _run_html(r) -> str:
    t = htmllib.escape(r.text)
    if r.bold:
        t = f"<b>{t}</b>"
    if r.italic:
        t = f"<i>{t}</i>"
    if r.underline:
        t = f"<u>{t}</u>"
    return t


def _row_cells(row):
    """(cell, colspan) pairs: python-docx repeats a merged cell once per grid column it covers."""
    out = []
    for c in row.cells:
        if out and out[-1][0]._tc is c._tc:
            out[-1][1] += 1
        else:
            out.append([c, 1])
    return out


def docx_to_txt(src: Path, out: Path, params: dict):
    import docx
    d = docx.Document(str(src))
    lines = []
    for block in d.iter_inner_content():  # paragraphs and tables in document order (tables used to go last)
        if isinstance(block, docx.table.Table):
            lines += ["\t".join(c.text for c, _ in _row_cells(row)) for row in block.rows]
        else:
            lines.append(block.text)
    out.write_text("\n".join(lines), encoding="utf-8")
    return [out]


def docx_to_pdf_fallback(src: Path, out: Path, params: dict):
    tmp = out.with_suffix(".tmp.html")
    docx_to_html(src, tmp, params)
    try:
        conv_pdf.html_to_pdf(tmp, out, params)
    finally:
        tmp.unlink(missing_ok=True)
    return [out]


def md_to_html(src: Path, out: Path, params: dict):
    import markdown
    body = markdown.markdown(read_text(src),
                             extensions=["tables", "fenced_code", "sane_lists", "toc"])
    out.write_text(f"<!doctype html><html><head><meta charset='utf-8'><title>{htmllib.escape(src.stem)}</title>"
                   f"</head><body>{body}</body></html>", encoding="utf-8")
    return [out]


def txt_to_html(src: Path, out: Path, params: dict):
    text = htmllib.escape(read_text(src))
    out.write_text(f"<!doctype html><meta charset='utf-8'><body><pre style='white-space:pre-wrap'>{text}</pre></body>",
                   encoding="utf-8")
    return [out]


def html_to_txt(src: Path, out: Path, params: dict):
    import re
    raw = read_text(src)
    raw = re.sub(r"(?is)<(script|style).*?</\1>", "", raw)
    raw = re.sub(r"(?i)<br\s*/?>|</p>|</h\d>|</li>|</tr>", "\n", raw)
    text = htmllib.unescape(re.sub(r"<[^>]+>", "", raw))
    out.write_text(re.sub(r"\n{3,}", "\n\n", text).strip(), encoding="utf-8")
    return [out]


def txt_to_docx(src: Path, out: Path, params: dict):
    import docx
    d = docx.Document()
    for line in read_text(src).splitlines():
        d.add_paragraph(line)
    d.save(str(out))
    return [out]


def md_to_docx(src: Path, out: Path, params: dict):
    import re
    import docx
    d = docx.Document()
    for line in read_text(src).splitlines():
        m = re.match(r"^(#{1,6})\s+(.*)", line)
        if m:
            d.add_heading(m.group(2), level=len(m.group(1)))
        elif re.match(r"^\s*[-*+]\s+", line):
            d.add_paragraph(re.sub(r"^\s*[-*+]\s+", "", line), style="List Bullet")
        elif re.match(r"^\s*\d+\.\s+", line):
            d.add_paragraph(re.sub(r"^\s*\d+\.\s+", "", line), style="List Number")
        elif line.strip():
            p = d.add_paragraph()
            for chunk in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", line):
                if chunk.startswith("**") and chunk.endswith("**"):
                    p.add_run(chunk[2:-2]).bold = True
                elif chunk.startswith("*") and chunk.endswith("*") and len(chunk) > 1:
                    p.add_run(chunk[1:-1]).italic = True
                else:
                    p.add_run(chunk)
    d.save(str(out))
    return [out]


def html_to_md(src: Path, out: Path, params: dict):
    import re
    raw = read_text(src)
    raw = re.sub(r"(?is)<(script|style|head).*?</\1>", "", raw)
    for n in range(6, 0, -1):
        raw = re.sub(rf"(?is)<h{n}[^>]*>(.*?)</h{n}>", lambda m: "\n" + "#" * n + " " + m.group(1) + "\n", raw)
    raw = re.sub(r"(?is)<(b|strong)[^>]*>(.*?)</\1>", r"**\2**", raw)
    raw = re.sub(r"(?is)<(i|em)[^>]*>(.*?)</\1>", r"*\2*", raw)
    raw = re.sub(r"(?is)<a[^>]*href=['\"]([^'\"]+)['\"][^>]*>(.*?)</a>", r"[\2](\1)", raw)
    raw = re.sub(r"(?is)<li[^>]*>", "\n- ", raw)
    raw = re.sub(r"(?i)<br\s*/?>|</p>|</tr>", "\n\n", raw)
    text = htmllib.unescape(re.sub(r"<[^>]+>", "", raw))
    out.write_text(re.sub(r"\n{3,}", "\n\n", text).strip() + "\n", encoding="utf-8")
    return [out]
