"""Tests for the Docs & sheets editor. Run:  .venv\\Scripts\\python.exe tests\\office\\test_office.py [--big]
(also works under pytest). Sample files are generated into a temp folder."""
import base64
import csv
import io
import os
import re
import sys
import tempfile
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(Path(__file__).parent))

import make_samples  # noqa: E402
from core import office_editor as oe  # noqa: E402

WORK = Path(os.environ.get("OFFICE_TEST_DIR") or tempfile.mkdtemp(prefix="office_test_"))
WORK.mkdir(parents=True, exist_ok=True)


def sample(name):
    p = WORK / name
    if not p.exists():
        {"rich.docx": make_samples.make_docx, "multi.xlsx": make_samples.make_xlsx, "tricky.csv": make_samples.make_csv,
         "big.xlsx": make_samples.make_big_xlsx}[name](p)
    return p


def strip_imgs(h):
    return re.sub(r'base64,[^"]+', "B64", h)


# ------------------------------------------------------------------ documents
def test_docx_to_html_features():
    h = oe.docx_to_html(sample("rich.docx"))["html"]
    for needle in ['<h1 class="title">Quarterly Report', "<h1>Introduction", "<b>bold</b>", "<i>italic</i>",
                   "<u>underlined</u>", "<s>struck</s>", "font-size:20pt", "color:#c00000", "background-color:#ffff00",
                   'href="https://example.com/page"', "text-align:center", "text-align:right", "text-align:justify",
                   "<ul><li>Bullet one", "<ul><li>Nested bullet", "<ol><li>Step one", 'colspan="2"', 'rowspan="2"',
                   "<blockquote>Quoted wisdom", "data:image/png;base64,", 'class="page-break"', "日本語"]:
        assert needle in h, needle


def test_docx_round_trip():
    src = sample("rich.docx")
    h1 = oe.docx_to_html(src)["html"]
    out = WORK / "roundtrip.docx"
    oe.html_to_docx(h1, out, template=src)
    h2 = oe.docx_to_html(out)["html"]
    assert strip_imgs(h1).replace(" </p>", "</p>") == strip_imgs(h2).replace(" </p>", "</p>")
    import docx
    d = docx.Document(str(out))
    styles = [p.style.name for p in d.paragraphs]
    assert "Title" in styles and "Heading 1" in styles and "Heading 2" in styles and "Quote" in styles
    assert len(d.tables) == 1 and d.tables[0].cell(0, 0).text == "Merged header"
    assert d.tables[0].cell(0, 0)._tc is d.tables[0].cell(0, 1)._tc  # still merged
    assert len(d.inline_shapes) == 1
    numbered = [p for p in d.paragraphs if p._p.pPr is not None and p._p.pPr.numPr is not None]
    assert len(numbered) == 6  # 4 bullets + 2 numbers, real Word numbering
    assert any(r.font.highlight_color for p in d.paragraphs for r in p.runs)
    xml = d.element.xml
    assert 'w:type="page"' in xml and "w:hyperlink" in xml


def test_html_to_docx_from_scratch_and_other_formats():
    png = base64.b64encode(make_samples.png_bytes()).decode()
    html = ('<h1 class="title">T</h1><h2>Sub</h2><p>Hello <b>b</b> <a href="https://x.org">x</a></p>'
            '<ul><li>a<ul><li>a1</li></ul></li><li>b</li></ul><ol start="3"><li>c</li></ol>'
            f'<p><img src="data:image/png;base64,{png}" style="width:50%"></p><hr><pre>code  line\nsecond</pre>'
            '<table><tr><th>H1</th><th>H2</th></tr><tr><td colspan="2">wide</td></tr></table>'
            '<div class="page-break" contenteditable="false"></div><p style="margin-left:36pt">indented</p>')
    out = WORK / "scratch.docx"
    oe.html_to_docx(html, out)
    back = oe.docx_to_html(out)["html"]
    for needle in ['<h1 class="title">T', "<h2>Sub", "<b>b</b>", 'href="https://x.org"', "<ul><li>a<ul><li>a1",
                   "<ol><li>c", "<img", "<hr>", "<pre>code  line\nsecond</pre>".replace("  ", " &nbsp;"),
                   'colspan="2"', 'class="page-break"', "margin-left:36pt"]:
        assert needle in back, (needle, back)
    md = oe.html_to_md(html)
    assert "# T" in md and "- a" in md and "    - a1" in md and "3. c" in md and "```" in md and "| H1 | H2 |" in md
    txt = oe.html_to_text(html)
    assert "Hello b x" in txt and "- a" in txt and "  - a1" in txt
    for fmt in ("html", "txt", "md", "pdf"):
        dest = WORK / f"scratch_out.{fmt}"
        oe.save_document("", html, fmt, dest)
        assert dest.exists() and dest.stat().st_size > 50, fmt


def test_open_txt_md_html():
    t = WORK / "a.txt"
    t.write_text("line one\n\tTabbed  two\n", encoding="utf-8")
    assert '<span class="tab">' in oe.open_document(t)["html"]
    m = WORK / "a.md"
    m.write_text("# Head\n\n- x\n- y\n\n| a | b |\n|---|---|\n| 1 | 2 |\n", encoding="utf-8")
    h = oe.open_document(m)["html"]
    assert "<h1>Head</h1>" in h and "<li>x</li>" in h and "<table>" in h
    w = WORK / "a.html"
    w.write_text('<html><head><style>body{}</style></head><body><p onclick="x()">Hi<script>bad()</script></p></body></html>')
    h = oe.open_document(w)["html"]
    assert "Hi" in h and "script" not in h and "onclick" not in h and "style" not in h


# ------------------------------------------------------------------ spreadsheets
def wait_ready(s, timeout=600):
    t = time.time()
    while s.status == "loading" and time.time() - t < timeout:
        time.sleep(0.1)
    assert s.status == "ready", s.error
    return s


def test_xlsx_load_edit_save():
    import openpyxl
    src = sample("multi.xlsx")
    s = wait_ready(oe.open_sheet(src))
    meta = s.meta()
    assert [x["name"] for x in meta["sheets"]] == ["Sales", "Notes"]
    assert meta["sheets"][0]["merged"] == [[8, 0, 8, 2]] and meta["sheets"][0]["frozen"] == 1
    rows = s.rows(0, 0, 10)["rows"]
    assert rows[1][3]["f"] == "=B2*C2"
    assert rows[1][2]["v"] == "$1.25"            # number format applied
    assert rows[1][4]["v"] == "2024-01-02"       # date format
    assert rows[6][3]["v"] in ("16.95", "=SUM(D2:D5)")  # cached (or evaluated) value
    # edit: value change recalculates the formulas we understand
    s.apply({"type": "set", "sheet": 0, "cells": [[1, 1, "10"]]})  # B2 = 10
    rows = s.rows(0, 0, 10)["rows"]
    assert rows[1][3]["v"] == "12.5" and rows[6][3]["v"] == "25.7"
    # unsupported function shows its formula text
    s.apply({"type": "set", "sheet": 0, "cells": [[10, 0, "=VLOOKUP(1,A1:B2,2)"], [11, 0, "=SUM(B2:B5)/COUNT(B2:B5)"]]})
    rows = s.rows(0, 10, 2)["rows"]
    assert rows[0][0]["v"] == "=VLOOKUP(1,A1:B2,2)" and rows[0][0].get("x")
    assert rows[1][0]["v"] == "7.25"
    # formatting + undo/redo
    s.apply({"type": "fmt", "sheet": 0, "r1": 1, "c1": 0, "r2": 2, "c2": 0, "patch": {"b": True, "fill": "#ffcc00"}})
    assert s.rows(0, 1, 1)["styles"][0].get("b") == 1
    s.do_undo()
    assert not any(st.get("b") for st in s.rows(0, 1, 1)["styles"])
    s.do_redo()
    assert any(st.get("b") for st in s.rows(0, 1, 1)["styles"])
    # insert a row above row 3: formulas below shift
    s.apply({"type": "insert", "sheet": 0, "axis": "row", "at": 2, "n": 1})
    assert s.wb["Sales"]["D8"].value == "=SUM(D2:D6)"
    assert s.wb["Notes"]["B2"].value == "=Sales!D8*2"
    s.do_undo()
    assert s.wb["Sales"]["D7"].value == "=SUM(D2:D5)" and s.wb["Notes"]["B2"].value == "=Sales!D7*2"
    # rename sheet updates references
    s.apply({"type": "rename_sheet", "sheet": 0, "name": "Sales 2024"})
    assert s.wb["Notes"]["B2"].value == "='Sales 2024'!D7*2"
    s.do_undo()
    assert s.wb["Notes"]["B2"].value == "=Sales!D7*2"
    out = WORK / "multi_edited.xlsx"
    s.write("xlsx", out)
    wb = openpyxl.load_workbook(out)
    ws = wb["Sales"]
    assert ws["B2"].value == 10 and ws["D2"].value == "=B2*C2" and ws["D7"].value == "=SUM(D2:D5)"
    assert ws["E7"].value.startswith("=VLOOKUP")
    assert ws["A2"].font.b and ws["A2"].fill.fgColor.rgb.endswith("FFCC00")
    assert ws["C3"].number_format.replace("\\", "") == "$#,##0.00" and ws["B1"].fill.fgColor.rgb.endswith("FFE3A3")
    assert "A9:C9" in [str(m) for m in ws.merged_cells.ranges] and ws.freeze_panes == "A2"
    assert wb["Notes"]["A1"].value == "Keep me" and wb["Notes"]["B2"].value == "=Sales!D7*2"
    assert s.meta()["dirty"] is False


def test_sort_and_delete_and_sheets():
    s = wait_ready(oe.open_sheet(sample("multi.xlsx")))
    with_merge = False
    try:
        s.apply({"type": "sort", "sheet": 0, "col": 1, "asc": True, "header": True})
    except ValueError:
        with_merge = True  # A9:C9 is merged -> refused
    assert with_merge
    s.apply({"type": "add_sheet", "name": "Data"})
    ws = s.wb["Data"]
    s.apply({"type": "set", "sheet": 2, "cells": [[0, 0, "n"], [1, 0, "3"], [2, 0, "1"], [3, 0, "2"], [4, 0, "x"]]})
    s.apply({"type": "sort", "sheet": 2, "col": 0, "asc": True, "header": True})
    assert [ws.cell(r, 1).value for r in range(1, 6)] == ["n", 1, 2, 3, "x"]
    s.apply({"type": "sort", "sheet": 2, "col": 0, "asc": False, "header": True})
    assert [ws.cell(r, 1).value for r in range(1, 6)] == ["n", "x", 3, 2, 1]
    s.do_undo()
    assert [ws.cell(r, 1).value for r in range(1, 6)] == ["n", 1, 2, 3, "x"]
    s.apply({"type": "delete", "sheet": 2, "axis": "row", "at": 1, "n": 2})
    assert [ws.cell(r, 1).value for r in range(1, 4)] == ["n", 3, "x"]
    s.do_undo()
    assert [ws.cell(r, 1).value for r in range(1, 6)] == ["n", 1, 2, 3, "x"]
    # deleting a column referenced by a formula -> #REF!, undo restores
    s.apply({"type": "delete", "sheet": 0, "axis": "col", "at": 1, "n": 1})
    assert s.wb["Sales"]["C2"].value == "=#REF!*B2"
    s.do_undo()
    assert s.wb["Sales"]["D2"].value == "=B2*C2"
    s.apply({"type": "delete_sheet", "sheet": 2})
    assert s.wb.sheetnames == ["Sales", "Notes"]
    s.do_undo()
    assert s.wb.sheetnames == ["Sales", "Notes", "Data"]
    s.apply({"type": "col_width", "sheet": 0, "col": 0, "px": 200})
    assert s.meta()["sheets"][0]["widths"][0] == 200
    s.apply({"type": "freeze", "sheet": 0, "rows": 0})
    assert s.meta()["sheets"][0]["frozen"] == 0
    assert s.search(0, "pear") == [2]


def test_csv_round_trip():
    src = sample("tricky.csv")
    s = wait_ready(oe.open_sheet(src))
    rows = s.rows(0, 0, 10)["rows"]
    assert rows[1][0]["v"] == "Zoë" and rows[1][1]["v"] == 'She said "hi", then left'
    assert rows[1][3]["v"] == "007" and "n" not in rows[1][3]  # text kept, not a number
    assert rows[3][1]["v"] == "line\nbreak"
    s.apply({"type": "set", "sheet": 0, "cells": [[2, 2, "4.75"], [4, 0, "new, row"]]})
    out = WORK / "tricky_out.csv"
    s.write("csv", out)
    with open(out, encoding="utf-8", newline="") as f:
        data = list(csv.reader(f))
    with open(src, encoding="utf-8", newline="") as f:
        orig = list(csv.reader(f))
    assert data[:2] == orig[:2] and data[3] == orig[3]
    assert data[2] == ["日本", "comma, inside", "4.75", "0042"] and data[4][0] == "new, row"
    for fmt in ("xlsx", "json", "ods", "tsv"):
        dest = WORK / f"tricky_out.{fmt}"
        s.write(fmt, dest)
        assert dest.stat().st_size > 20


def test_number_formats():
    import datetime as dt
    assert oe.fmt_value(1234.5, "#,##0.00") == "1,234.50"
    assert oe.fmt_value(0.256, "0.0%") == "25.6%"
    assert oe.fmt_value(-3, '#,##0;(#,##0)') == "(3)"
    assert oe.fmt_value(5, '[$€-407] #,##0.00') == "€ 5.00"
    assert oe.fmt_value(dt.datetime(2024, 3, 5, 14, 7), "dd/mm/yyyy hh:mm") == "05/03/2024 14:07"
    assert oe.fmt_value(dt.datetime(2024, 3, 5), "mmm d, yyyy") == "Mar 5, 2024"
    assert oe.fmt_value(0.1 + 0.2) == "0.3"
    assert oe._parse_input("007") == ("007", "s", None) and oe._parse_input("1e3")[0] == 1000.0


def test_big_paging(n_check=True):
    src = sample("big.xlsx")
    t = time.time()
    s = oe.open_sheet(src, wait=0.5)
    first = s.meta()
    assert first["status"] in ("loading", "ready")
    if first["status"] == "loading":
        assert first["preview"] and len(first["preview"]["rows"]) == 200
    wait_ready(s)
    load = time.time() - t
    meta = s.meta()
    assert meta["sheets"][0]["rows"] == 100_001
    t = time.time()
    for start in (0, 50_000, 99_900):
        page = s.rows(0, start, 200)
        assert page["rows"][0][0]["v"] == ("id" if start == 0 else str(start))
    page_ms = (time.time() - t) / 3 * 1000
    t = time.time()
    hits = s.search(0, "kolkata")
    search_ms = (time.time() - t) * 1000
    assert len(hits) == 20_000
    t = time.time()
    s.apply({"type": "set", "sheet": 0, "cells": [[5, 1, "edited"]]})
    edit_ms = (time.time() - t) * 1000
    print(f"  big.xlsx: load {load:.1f}s, page {page_ms:.0f} ms, search {search_ms:.0f} ms, edit {edit_ms:.0f} ms")
    assert page_ms < 500 and edit_ms < 2000


# ------------------------------------------------------------------ HTTP endpoints
def client():
    import app
    return app.app.test_client()


def post(c, url, body):
    r = c.post(url, json=body, headers={"X-Local-App": "1"}, base_url="http://127.0.0.1")
    return r.get_json()


def get(c, url):
    return c.get(url, base_url="http://127.0.0.1").get_json()


def test_endpoints():
    c = client()
    assert get(c, "/api/office/caps")["ok"]
    # no header -> refused
    assert c.post("/api/office/open", json={"path": "x"}, base_url="http://127.0.0.1").status_code == 403
    out_dir = WORK / "out"
    d = post(c, "/api/office/open", {"path": str(sample("rich.docx"))})
    assert d["ok"] and d["kind"] == "doc" and "<table>" in d["html"]
    html = d["html"].replace("Introduction", "Intro (edited)")
    r1 = post(c, "/api/office/doc/save", {"path": d["path"], "html": html, "format": "docx", "out_dir": str(out_dir)})
    assert r1["ok"] and r1["name"].startswith("rich (edited)")
    r2 = post(c, "/api/office/doc/save", {"path": d["path"], "html": html, "format": "docx", "out_dir": str(out_dir)})
    assert r2["path"] == r1["path"]  # second Save updates the same copy
    assert "Intro (edited)" in oe.docx_to_html(r1["path"])["html"]
    for fmt in ("pdf", "html", "md", "txt"):
        r = post(c, "/api/office/doc/save", {"path": d["path"], "html": html, "format": fmt, "out_dir": str(out_dir)})
        assert r["ok"] and Path(r["path"]).exists(), (fmt, r)
    bad = post(c, "/api/office/doc/save", {"path": d["path"], "html": html, "format": "pdf", "mode": "overwrite", "out_dir": str(out_dir)})
    assert not bad["ok"]
    copy = WORK / "overwrite_me.docx"
    copy.write_bytes(sample("rich.docx").read_bytes())
    r = post(c, "/api/office/doc/save", {"path": str(copy), "html": html, "format": "docx", "mode": "overwrite"})
    assert r["ok"] and r["path"] == str(copy) and "Intro (edited)" in oe.docx_to_html(copy)["html"]

    d = post(c, "/api/office/open", {"path": str(sample("multi.xlsx"))})
    assert d["ok"] and d["kind"] == "sheet"
    sid = d["sheet"]["sid"]
    while d["sheet"]["status"] == "loading":
        time.sleep(0.2)
        d = {"sheet": get(c, f"/api/office/sheet/meta?sid={sid}")["sheet"]}
    rows = get(c, f"/api/office/sheet/rows?sid={sid}&sheet=0&start=0&count=50")
    assert rows["ok"] and rows["rows"][0][0]["v"] == "Item"
    assert post(c, "/api/office/sheet/rows", {"sid": sid, "sheet": 0, "indices": [3, 1]})["rows"][0][0]["v"] == "Plums"
    r = post(c, "/api/office/sheet/op", {"sid": sid, "op": {"type": "set", "sheet": 0, "cells": [[1, 0, "Kiwi"]]}})
    assert r["ok"] and r["sheet"]["dirty"] and r["sheet"]["can_undo"]
    assert not post(c, "/api/office/sheet/op", {"sid": sid, "op": {"type": "_restore"}})["ok"]
    assert post(c, "/api/office/sheet/undo", {"sid": sid})["sheet"]["can_redo"]
    assert post(c, "/api/office/sheet/redo", {"sid": sid})["ok"]
    assert post(c, "/api/office/sheet/search", {"sid": sid, "sheet": 0, "q": "kiwi"})["rows"] == [1]
    rng = post(c, "/api/office/sheet/range", {"sid": sid, "sheet": 0, "r1": 0, "c1": 0, "r2": 1, "c2": 3, "raw": True})
    assert rng["values"][1] == ["Kiwi", "3", "1.25", "=B2*C2"]
    pr = post(c, "/api/office/sheet/print", {"sid": sid, "sheet": 0})
    assert "<table>" in pr["html"]
    sv = post(c, "/api/office/sheet/save", {"sid": sid, "format": "xlsx", "out_dir": str(out_dir)})
    assert sv["ok"] and sv["name"].startswith("multi (edited)") and not sv["sheet"]["dirty"]
    for fmt in ("csv", "json", "ods", "pdf"):
        r = post(c, "/api/office/sheet/save", {"sid": sid, "format": fmt, "out_dir": str(out_dir)})
        assert r["ok"] and Path(r["path"]).exists(), (fmt, r)
    assert not post(c, "/api/office/sheet/save", {"sid": sid, "format": "csv", "mode": "overwrite"})["ok"]
    assert post(c, "/api/office/sheet/close", {"sid": sid})["ok"]
    assert not get(c, f"/api/office/sheet/meta?sid={sid}")["ok"]

    d = post(c, "/api/office/open", {"path": str(sample("tricky.csv"))})
    assert d["ok"] and d["kind"] == "sheet"
    pdfp = WORK / "out" / [p for p in os.listdir(WORK / "out") if p.endswith(".pdf")][0]
    d = post(c, "/api/office/open", {"path": str(pdfp)})
    assert d["ok"] and d["kind"] == "pdf" and d["pages"]
    pp = WORK / "deck.pptx"
    if pp.exists():
        d = post(c, "/api/office/open", {"path": str(pp)})
        assert d["ok"] and d["kind"] == "slides", d
    assert not post(c, "/api/office/open", {"path": str(WORK / "nope.docx")})["ok"]
    nd = post(c, "/api/office/new", {"kind": "doc"})
    assert nd["ok"] and nd["file"]["temp"]
    r = post(c, "/api/office/doc/save", {"path": nd["file"]["path"], "html": "<p>Hi</p>", "format": "docx", "out_dir": str(out_dir)})
    assert r["ok"] and r["name"].startswith("New document") and "(edited)" not in r["name"]
    assert not post(c, "/api/office/doc/save", {"path": nd["file"]["path"], "html": "<p>x</p>", "format": "docx", "mode": "overwrite"})["ok"]


if __name__ == "__main__":
    tests = [v for k, v in list(globals().items()) if k.startswith("test_") and callable(v)]
    if "--big" not in sys.argv:
        tests = [t for t in tests if t.__name__ != "test_big_paging"]
    failed = 0
    for t in tests:
        t0 = time.time()
        try:
            t()
            print(f"PASS {t.__name__} ({time.time() - t0:.1f}s)")
        except Exception as e:  # noqa
            failed += 1
            import traceback
            traceback.print_exc()
            print(f"FAIL {t.__name__}: {e!r}")
    print(f"\n{len(tests) - failed}/{len(tests)} passed  (files in {WORK})")
    sys.exit(1 if failed else 0)
