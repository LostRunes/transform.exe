"""Builds sample files for the Docs & sheets editor tests.
Usage: python tests/office/make_samples.py <out_dir>"""
import csv
import datetime as dt
import io
import sys
from pathlib import Path


def png_bytes(color=(240, 120, 160), size=(120, 80)):
    from PIL import Image, ImageDraw
    im = Image.new("RGB", size, color)
    ImageDraw.Draw(im).ellipse((10, 10, size[0] - 10, size[1] - 10), fill=(255, 255, 255))
    buf = io.BytesIO()
    im.save(buf, "PNG")
    return buf.getvalue()


def add_link(paragraph, text, url):
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    rid = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hl = OxmlElement("w:hyperlink")
    hl.set(qn("r:id"), rid)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(u)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    r.append(t)
    hl.append(r)
    paragraph._p.append(hl)


def make_docx(path: Path):
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
    from docx.shared import Inches, Pt, RGBColor
    d = docx.Document()
    d.add_heading("Quarterly Report", 0)  # Title
    d.add_heading("Introduction", 1)
    p = d.add_paragraph("Plain start, ")
    p.add_run("bold").bold = True
    p.add_run(", ")
    p.add_run("italic").italic = True
    p.add_run(", ")
    p.add_run("underlined").underline = True
    p.add_run(", ")
    r = p.add_run("struck")
    r.font.strike = True
    p.add_run(", ")
    r = p.add_run("big red")
    r.font.size = Pt(20)
    r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    p.add_run(", ")
    r = p.add_run("highlighted")
    r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    p.add_run(" and a link: ")
    add_link(p, "example site", "https://example.com/page")
    p.add_run(".")
    c = d.add_paragraph("Centered paragraph")
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = d.add_paragraph("Right aligned")
    rr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    j = d.add_paragraph("Justified text " * 12)
    j.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    d.add_heading("Lists", 2)
    d.add_paragraph("Bullet one", style="List Bullet")
    d.add_paragraph("Bullet two", style="List Bullet")
    d.add_paragraph("Nested bullet", style="List Bullet 2")
    d.add_paragraph("Bullet three", style="List Bullet")
    d.add_paragraph("Step one", style="List Number")
    d.add_paragraph("Step two", style="List Number")
    d.add_heading("A table", 3)
    t = d.add_table(rows=3, cols=3)
    t.style = d.styles["Table Grid"]
    for ri in range(3):
        for ci in range(3):
            t.cell(ri, ci).text = f"R{ri + 1}C{ci + 1}"
    a = t.cell(0, 0).merge(t.cell(0, 1))
    a.text = "Merged header"
    b = t.cell(1, 2).merge(t.cell(2, 2))
    b.text = "Tall"
    d.add_paragraph("Quoted wisdom", style="Quote")
    d.add_paragraph("An image:")
    d.add_picture(io.BytesIO(png_bytes()), width=Inches(1.5))
    pb = d.add_paragraph()
    pb.add_run().add_break(WD_BREAK.PAGE)
    d.add_heading("Second page", 1)
    d.add_paragraph("Unicode: café, naïve, 日本語, emoji ✓")
    d.save(str(path))


def make_xlsx(path: Path):
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Item", "Qty", "Price", "Total", "Date"])
    rows = [("Apples", 3, 1.25), ("Pears", 5, 0.8), ("Plums", 12, 0.35), ("Figs", 2, 2.5)]
    for i, (n, q, pr) in enumerate(rows, start=2):
        ws.append([n, q, pr, f"=B{i}*C{i}", dt.date(2024, 1, i)])
        ws[f"C{i}"].number_format = "$#,##0.00"
        ws[f"E{i}"].number_format = "yyyy-mm-dd"
    ws["A7"] = "Sum"
    ws["D7"] = "=SUM(D2:D5)"
    ws["E7"] = "=VLOOKUP(\"Figs\",A2:D5,4,FALSE)"
    ws["A1"].font = Font(bold=True)
    ws["B1"].fill = PatternFill("solid", fgColor="FFE3A3")
    ws.merge_cells("A9:C9")
    ws["A9"] = "Merged note across three columns"
    ws["F2"] = 0.256
    ws["F2"].number_format = "0.0%"
    ws.freeze_panes = "A2"
    ws2 = wb.create_sheet("Notes")
    ws2["A1"] = "Keep me"
    ws2["B2"] = "=Sales!D7*2"
    wb.save(path)
    # cached values: let Excel compute them if available, so data_only loads see numbers
    try:
        from core import conv_doc
        if conv_doc.HAS_EXCEL:
            def run():
                import win32com.client
                xl = win32com.client.DispatchEx("Excel.Application")
                xl.Visible = False
                xl.DisplayAlerts = False
                try:
                    b = xl.Workbooks.Open(str(path.resolve()))
                    b.Save()
                    b.Close(False)
                finally:
                    xl.Quit()
            conv_doc._com(run)
    except Exception as e:  # noqa
        print("Excel recalculation skipped:", e)


def make_big_xlsx(path: Path, n=100_000):
    import openpyxl
    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet("Big")
    ws.append(["id", "name", "value", "city", "group"])
    cities = ["Pune", "Delhi", "Kolkata", "Bhubaneswar", "Mumbai"]
    for i in range(1, n + 1):
        ws.append([i, f"name {i}", i * 1.5, cities[i % 5], i % 7])
    wb.save(path)


def make_csv(path: Path):
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["name", "quote", "amount", "code"])
        w.writerow(["Zoë", 'She said "hi", then left', "12.50", "007"])
        w.writerow(["日本", "comma, inside", "3", "0042"])
        w.writerow(["Ünïcode ✓", "line\nbreak", "-4.25", "x"])


def make_pptx(path: Path):
    """Needs PowerPoint (python-pptx isn't a dependency)."""
    from core import conv_doc
    if not conv_doc.HAS_POWERPOINT:
        return False

    def run():
        import win32com.client
        pp = win32com.client.DispatchEx("PowerPoint.Application")
        try:
            pres = pp.Presentations.Add(WithWindow=False)
            for i in range(1, 4):
                s = pres.Slides.Add(i, 2)  # title + text
                s.Shapes(1).TextFrame.TextRange.Text = f"Slide {i}: hello"
                s.Shapes(2).TextFrame.TextRange.Text = "First point\rSecond point"
            pres.SaveAs(str(path.resolve()), 24)
            pres.Close()
        finally:
            pp.Quit()
    conv_doc._com(run)
    return True


if __name__ == "__main__":
    out = Path(sys.argv[1])
    out.mkdir(parents=True, exist_ok=True)
    sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
    make_docx(out / "rich.docx")
    make_xlsx(out / "multi.xlsx")
    make_csv(out / "tricky.csv")
    if "--big" in sys.argv:
        make_big_xlsx(out / "big.xlsx")
    print("samples in", out)
